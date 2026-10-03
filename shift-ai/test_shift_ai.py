"""既存Excelと同じ構造の小さなブックを作って、学習 → 作成 → 書き込みを通しで確かめる"""

import zipfile

import openpyxl
import pytest

import shift_ai as S


def make_book(path, grid):
    wb = openpyxl.Workbook()
    wb.active.title = "メインメニュー"
    ws = wb.create_sheet(S.SHEET_SHIFT)
    # NO, 始業, 終業, 休1入, 休1出, -, -, 就業計
    ws.append([])
    ws.append(["ｼﾌﾄNO.", "始業", "終業", "休１入", "休１出", "休２入", "休２出", "就業計"])
    ws.append([1, 945, 2100, 1500, 1630, None, None, 9.75])
    ws.append([7, 1100, 1500, None, None, None, None, 4])
    ws.append([22, 1900, 2100, None, None, None, None, 2])
    ws.append([801, "オペレーション外勤務", None, None, None, None, None, 8])

    ws = wb.create_sheet(S.SHEET_PERSONAL)
    ws["B1"], ws["E1"] = 2025, 11
    for d in range(30):
        ws.cell(2, S.PERSONAL_DAY1_COL + d, d + 1)
    for i, (code, name, row) in enumerate(grid):
        r = S.PERSONAL_FIRST_ROW + i
        ws.cell(r, 1, code)
        ws.cell(r, 2, i + 1)
        ws.cell(r, 3, name)
        for d, v in enumerate(row):
            if v is not None:
                ws.cell(r, S.PERSONAL_DAY1_COL + d, v)
    r = S.PERSONAL_FIRST_ROW + len(grid)
    ws.cell(r, 3, "(祝日は1）→")
    ws.cell(r, S.PERSONAL_DAY1_COL + 2, 1)          # 11/3 祝日

    ws = wb.create_sheet(S.SHEET_MODEL)
    ws.append([])
    ws.append([])
    ws.append([])
    ws.append(["パターン"] + [f"～{h + 1}" for h in range(7, 31)])
    for label, people in [("平日①", 1), ("平日②", 1), ("土曜", 2), ("日祝", 2)]:
        row = [label] + [0] * 24
        for h in range(10, 21):
            row[1 + h - S.MODEL_FIRST_HOUR] = people
        ws.append(row)
    wb.save(path)


def test_shift_minutes():
    sh = S.Shift(1, 9 * 60 + 45, 21 * 60, [(15 * 60, 16 * 60 + 30)], 9.75)
    assert sh.minutes_in_hour(9) == 15
    assert sh.minutes_in_hour(10) == 60
    assert sh.minutes_in_hour(15) == 0
    assert sh.minutes_in_hour(16) == 30
    assert sh.minutes_in_hour(21) == 0
    assert S.hhmm_to_min(2530) == 25 * 60 + 30


def test_learn_and_generate(tmp_path):
    full = [1 if d % 7 != 2 else 0 for d in range(30)]           # 月曜休み
    part = [7 if d % 7 in (0, 1) else None for d in range(30)]    # 土日だけ
    src = tmp_path / "src.xlsx"
    make_book(src, [(1, "社員A", full), (5, "パートB", part), (6, "パートC", [None] * 30)])

    book = S.load_workbook(str(src))
    assert book.num_days == 30 and book.holidays == {3}
    assert [s.name for s in book.staff] == ["社員A", "パートB", "パートC"]
    assert book.demand["土曜"][10] == 2

    rules = S.learn(book)
    assert rules["社員A"].allowed_shifts == [1]
    assert "月" not in rules["社員A"].allowed_days
    assert rules["パートB"].allowed_days == "土日"
    assert rules["パートB"].off_mark == "空欄"
    assert not rules["パートC"].target

    settings = tmp_path / "AI設定.xlsx"
    S.save_settings(str(settings), book, rules)
    rules2, glob = S.load_settings(str(settings))
    assert rules2["パートB"].allowed_shifts == [7]
    glob["計算時間(秒)"] = 10

    # 希望: 社員Aは11/5休み、パートBは11/8にシフト22
    book.staff[0].grid = [None] * 30
    book.staff[0].grid[4] = 0
    book.staff[1].grid = [None] * 30
    book.staff[1].grid[7] = 22
    result = S.solve(book, rules2, glob)
    assert result.assignment, result.status
    assert result.assignment[0, 5] == 0
    assert result.assignment[1, 8] == 22
    a = [result.assignment[0, d] for d in range(1, 31)]
    consec = rules2["社員A"].max_consecutive
    assert consec == 6                                           # 学習した連勤（6勤1休）
    for d0 in range(30 - consec):
        assert sum(1 for v in a[d0:d0 + consec + 1] if v) <= consec
    for d in range(1, 31):
        if book.weekday(d) not in "土日" and d != 8:
            assert result.assignment[1, d] is None

    out = tmp_path / "out.xlsx"
    S.write_assignment(book, result, str(out))
    assert zipfile.ZipFile(out).testzip() is None
    ws = openpyxl.load_workbook(out)[S.SHEET_PERSONAL]
    assert ws.cell(5, S.PERSONAL_DAY1_COL + 4).value == 0
    assert ws.cell(6, S.PERSONAL_DAY1_COL + 7).value == 22
    assert ws.cell(5, 3).value == "社員A"
    assert ws.cell(2, S.PERSONAL_DAY1_COL).value == 1


def test_set_cells_inserts_in_column_order():
    xml = '<sheetData><row r="5"><c r="C5" t="s"><v>0</v></c><c r="G5" s="3"><v>9</v></c></row><row r="6"/></sheetData>'
    out = S._set_cells(xml, {"E5": 1, "G5": None, "F6": 4})
    assert '<c r="C5" t="s"><v>0</v></c><c r="E5"><v>1</v></c><c r="G5" s="3"/>' in out
    assert '<row r="6"><c r="F6"><v>4</v></c></row>' in out


if __name__ == "__main__":
    pytest.main([__file__, "-q"])
