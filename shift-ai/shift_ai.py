"""AIシフト作成ツール

既存のシフト表Excel（.xlsm）の「個人シフト入力表」を、OR-Tools（CP-SAT）で自動的に埋める。

使い方:
    # 1) 完成済みの月のファイルから、スタッフごとの勤務傾向を学習して設定ファイルを作る
    python shift_ai.py learn  202511スケジュール.xlsm  -s AI設定.xlsx

    # 2) 設定ファイル（必要なら手直しする）を使って、新しい月のシフトを作る
    python shift_ai.py generate 202512スケジュール.xlsm -s AI設定.xlsx -o AIシフト完成.xlsm

「個人シフト入力表」に入っている値は希望として扱う:
    シフトNO.   → その日はそのシフトで確定
    0           → 休み希望（確定）
    空欄        → AIが決める
"""

from __future__ import annotations

import argparse
import calendar
import datetime
import os
import re
import sys
import warnings
import zipfile
from collections import Counter
from dataclasses import dataclass, field

import openpyxl
from ortools.sat.python import cp_model

warnings.filterwarnings("ignore", module="openpyxl")

# ==============================
# Excelの構造（既存ファイルに合わせる）
# ==============================
SHEET_SHIFT = "ｼﾌﾄ条件設定表"
SHEET_PERSONAL = "個人シフト入力表"
SHEET_MODEL = "モデル時間"
SHEET_MASTER = "社員マスタ登録"

SHIFT_FIRST_ROW = 3          # ｼﾌﾄ条件設定表: A=NO, B=始業, C=終業, D/E=休1, F/G=休2, H=就業計
NO_TIME_SHIFT_HOURS = 8      # 801(オペレーション外勤務) / 802(ヘルプ) は時刻なし

PERSONAL_YEAR = "B1"
PERSONAL_MONTH = "E1"
PERSONAL_FIRST_ROW = 5       # 個人シフト入力表: A=社員CD, B=NO, C=氏名, E列〜=1日〜
PERSONAL_LAST_ROW = 37
PERSONAL_DAY1_COL = 5        # E列

MODEL_FIRST_COL = 2          # モデル時間: B列 = 7時台
MODEL_FIRST_HOUR = 7

WEEKDAYS = "月火水木金土日"
HOLIDAY = "祝"

SETTINGS_STAFF = "スタッフ設定"
SETTINGS_GLOBAL = "全体設定"
STAFF_HEADERS = [
    "社員CD", "氏名", "区分", "対象(○/×)", "勤務可能シフト", "勤務可能曜日",
    "目標出勤日数", "最大出勤日数", "最大時間/月", "最大連勤", "休日の書き方(0/空欄)",
    "よく使うシフト(参考)",
]
DEFAULT_GLOBAL = {
    "計算時間(秒)": 60,
    "最大連勤(既定)": 5,
    "月": "平日①", "火": "平日①", "水": "平日①", "木": "平日①", "金": "平日①",
    "土": "土曜", "日": "日祝", "祝": "日祝",
    "重み:人手不足(15分あたり)": 10,
    "重み:人手過剰(15分あたり)": 2,
    "重み:目標日数とのずれ(1日あたり)": 8,
    "重み:いつもと違うシフト": 3,
}


# ==============================
# データ
# ==============================
@dataclass
class Shift:
    no: int
    start: int | None            # 分（7:00 = 420、翌1:00 = 1500）
    end: int | None
    breaks: list[tuple[int, int]]
    hours: float

    def minutes_in_hour(self, hour: int) -> int:
        """hour時台（hour:00〜hour+1:00）に実際に働く分数（休憩を除く）"""
        if self.start is None:
            return 0
        lo, hi = hour * 60, hour * 60 + 60
        m = max(0, min(hi, self.end) - max(lo, self.start))
        for bs, be in self.breaks:
            m -= max(0, min(hi, be, self.end) - max(lo, bs, self.start))
        return max(0, m)


@dataclass
class Staff:
    row: int
    code: int | None
    name: str
    kubun: str = ""
    grid: list = field(default_factory=list)   # 日ごとの入力値（None / 0 / シフトNO.）


@dataclass
class StaffRule:
    target: bool = True
    allowed_shifts: list[int] = field(default_factory=list)
    allowed_days: str = WEEKDAYS + HOLIDAY
    target_days: int | None = None
    max_days: int | None = None
    max_hours: float | None = None
    max_consecutive: int | None = None
    off_mark: str = "0"
    usual: dict[int, int] = field(default_factory=dict)


@dataclass
class Workbook:
    path: str
    year: int
    month: int
    num_days: int
    shifts: dict[int, Shift]
    staff: list[Staff]
    holidays: set[int]
    demand: dict[str, dict[int, float]]   # パターン名 → {時: 必要人時}
    kubun_names: dict[int, str]

    def day_category(self, day: int) -> str:
        if day in self.holidays:
            return HOLIDAY
        return WEEKDAYS[datetime.date(self.year, self.month, day).weekday()]

    def weekday(self, day: int) -> str:
        return WEEKDAYS[datetime.date(self.year, self.month, day).weekday()]


# ==============================
# 読み込み
# ==============================
def hhmm_to_min(v) -> int | None:
    """945 / "9:45" / time / 9.75 などを分に変換"""
    if v is None or v == "":
        return None
    if isinstance(v, (datetime.time, datetime.datetime)):
        return v.hour * 60 + v.minute
    if isinstance(v, str):
        v = v.strip()
        if ":" in v:
            h, m = v.split(":")[:2]
            return int(h) * 60 + int(m)
        if not v.replace(".", "", 1).isdigit():
            return None
        v = float(v)
    v = float(v)
    if v >= 100:                      # HHMM 形式（945, 2100, 2530）
        return int(v) // 100 * 60 + int(v) % 100
    return round(v * 60)              # 時間の小数（9.75）


def to_int(v) -> int | None:
    if v is None or v == "":
        return None
    try:
        return int(float(v))
    except (TypeError, ValueError):
        return None


def load_workbook(path: str) -> Workbook:
    wb = openpyxl.load_workbook(path, data_only=True, keep_vba=False)

    # --- シフト定義 ---
    shifts: dict[int, Shift] = {}
    ws = wb[SHEET_SHIFT]
    for r in range(SHIFT_FIRST_ROW, ws.max_row + 1):
        no = to_int(ws.cell(r, 1).value)
        if not no:
            continue
        start, end = hhmm_to_min(ws.cell(r, 2).value), hhmm_to_min(ws.cell(r, 3).value)
        if start is not None and end is not None:
            if end <= start:
                end += 24 * 60
            breaks = []
            for c in (4, 6):
                bs, be = hhmm_to_min(ws.cell(r, c).value), hhmm_to_min(ws.cell(r, c + 1).value)
                if bs is not None and be is not None and be > bs:
                    breaks.append((bs, be))
            hours = ws.cell(r, 8).value
            if not isinstance(hours, (int, float)) or hours <= 0:
                hours = (end - start - sum(e - s for s, e in breaks)) / 60
            shifts[no] = Shift(no, start, end, breaks, float(hours))
        elif no >= 800:               # 時刻を持たない特別シフト
            hours = ws.cell(r, 8).value
            shifts[no] = Shift(no, None, None, [], float(hours or NO_TIME_SHIFT_HOURS))

    # --- 年月・日数 ---
    ws = wb[SHEET_PERSONAL]
    year, month = to_int(ws[PERSONAL_YEAR].value), to_int(ws[PERSONAL_MONTH].value)
    if not year or not month:
        raise ValueError(f"「{SHEET_PERSONAL}」の {PERSONAL_YEAR}/{PERSONAL_MONTH} に年月がありません")
    num_days = calendar.monthrange(year, month)[1]

    # --- 社員区分 ---
    kubun_names: dict[int, str] = {}
    if SHEET_MASTER in wb.sheetnames:
        wm = wb[SHEET_MASTER]
        for r in range(2, wm.max_row + 1):
            code = to_int(wm.cell(r, 1).value)
            if code is not None:
                kubun_names[code] = str(wm.cell(r, 4).value or "")

    # --- スタッフ・祝日 ---
    staff: list[Staff] = []
    holidays: set[int] = set()
    for r in range(PERSONAL_FIRST_ROW, PERSONAL_LAST_ROW + 1):
        name = ws.cell(r, 3).value
        name = str(name).strip() if name is not None else ""
        cells = [ws.cell(r, PERSONAL_DAY1_COL + d).value for d in range(num_days)]
        if "祝日" in name:
            holidays = {d + 1 for d, v in enumerate(cells) if to_int(v) == 1}
            continue
        if not name or name == "不足":
            continue
        code = to_int(ws.cell(r, 1).value)
        grid = [to_int(v) if v not in (None, "", " ") else None for v in cells]
        staff.append(Staff(r, code, name, kubun_names.get(code, ""), grid))

    # --- モデル時間（必要人時） ---
    demand: dict[str, dict[int, float]] = {}
    ws = wb[SHEET_MODEL]
    for r in range(1, ws.max_row + 1):
        label = ws.cell(r, 1).value
        if not isinstance(label, str) or label.strip() in ("", "パターン"):
            continue
        hours = {}
        for i in range(24):
            v = ws.cell(r, MODEL_FIRST_COL + i).value
            if isinstance(v, (int, float)) and v > 0:
                hours[MODEL_FIRST_HOUR + i] = float(v)
        demand[label.strip()] = hours

    return Workbook(path, year, month, num_days, shifts, staff, holidays, demand, kubun_names)


# ==============================
# 設定ファイル（AI設定.xlsx）
# ==============================
def _fmt_list(xs) -> str:
    return ",".join(str(x) for x in xs)


def _parse_list(v) -> list[int]:
    if v is None:
        return []
    return [int(x) for x in re.findall(r"\d+", str(v))]


def learn(book: Workbook) -> dict[str, StaffRule]:
    """完成済みの月の「個人シフト入力表」から、スタッフごとの勤務傾向を学習する"""
    rules: dict[str, StaffRule] = {}
    for s in book.staff:
        worked = [(d + 1, v) for d, v in enumerate(s.grid) if v]
        counts = Counter(v for _, v in worked if v in book.shifts)
        days = {book.day_category(d) for d, _ in worked}
        if "日" in days and not any(d in book.holidays for d in range(1, book.num_days + 1)):
            days.add(HOLIDAY)
        run = best = 0
        for v in s.grid:
            run = run + 1 if v else 0
            best = max(best, run)
        rules[s.name] = StaffRule(
            target=bool(worked),
            allowed_shifts=sorted(k for k in counts if book.shifts[k].start is not None),
            allowed_days="".join(c for c in WEEKDAYS + HOLIDAY if c in days),
            target_days=len(worked),
            max_days=min(book.num_days, len(worked) + 2) if worked else None,
            max_hours=None,
            max_consecutive=max(best, 1) if worked else None,
            off_mark="0" if 0 in s.grid else "空欄",
            usual=dict(counts.most_common()),
        )
        # 1か月休みなしだった人は連勤の傾向が分からないので、全体設定の既定値に任せる
        if rules[s.name].max_consecutive == book.num_days:
            rules[s.name].max_consecutive = None
    return rules


def save_settings(path: str, book: Workbook, rules: dict[str, StaffRule], glob: dict | None = None):
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SETTINGS_STAFF
    ws.append(STAFF_HEADERS)
    for s in book.staff:
        r = rules.get(s.name, StaffRule(target=False))
        ws.append([
            s.code, s.name, s.kubun, "○" if r.target else "×", _fmt_list(r.allowed_shifts),
            r.allowed_days, r.target_days, r.max_days, r.max_hours, r.max_consecutive, r.off_mark,
            ",".join(f"{k}:{n}" for k, n in r.usual.items()),
        ])
    head = PatternFill("solid", fgColor="DDEBF7")
    for c in ws[1]:
        c.font, c.fill = Font(bold=True), head
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for col, w in zip("ABCDEFGHIJKL", [8, 16, 14, 9, 18, 14, 10, 10, 10, 8, 12, 22]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "C2"

    note = ws.max_row + 2
    for i, t in enumerate([
        "【書き方】",
        "対象: ○ = AIが割り当てる / × = AIは触らない（手入力のまま。人数には数える）",
        "勤務可能シフト: シフトNO.をカンマ区切り（例: 1,4,6）。空欄なら全シフト可",
        "勤務可能曜日: 月火水木金土日祝 から働ける曜日だけ書く（祝 = 個人シフト入力表で祝日に1がある日）",
        "目標出勤日数: この日数に近づける / 最大出勤日数・最大時間/月・最大連勤: これを超えない（空欄なら制限なし・既定値）",
        "休日の書き方: 休みの日に 0 を書くか、空欄のままにするか",
    ]):
        ws.cell(note + i, 1, t).font = Font(bold=(i == 0), color="555555")

    wg = wb.create_sheet(SETTINGS_GLOBAL)
    wg.append(["項目", "値", "説明"])
    desc = {
        "計算時間(秒)": "長くするほど良い答えを探す",
        "最大連勤(既定)": "スタッフ設定で空欄の人に使う",
        "月": "その曜日に使うモデル時間のパターン名", "祝": "祝日に使うパターン名",
        "重み:人手不足(15分あたり)": "大きいほど不足を嫌う",
        "重み:人手過剰(15分あたり)": "大きいほど人の余りを嫌う",
        "重み:目標日数とのずれ(1日あたり)": "大きいほど各自の目標日数を守る",
        "重み:いつもと違うシフト": "大きいほど普段のシフトを優先",
    }
    for k, v in (glob or DEFAULT_GLOBAL).items():
        wg.append([k, v, desc.get(k, "")])
    for c in wg[1]:
        c.font, c.fill = Font(bold=True), head
    wg.column_dimensions["A"].width = 30
    wg.column_dimensions["B"].width = 10
    wg.column_dimensions["C"].width = 36
    wb.save(path)


def load_settings(path: str) -> tuple[dict[str, StaffRule], dict]:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[SETTINGS_STAFF]
    head = [str(c.value or "") for c in ws[1]]
    idx = {h: i for i, h in enumerate(head)}

    def num(row, key, cast=int):
        v = row[idx[key]] if key in idx else None
        try:
            return cast(v) if v not in (None, "") else None
        except (TypeError, ValueError):
            return None

    rules: dict[str, StaffRule] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        name = row[idx["氏名"]]
        if not name or str(name).startswith("【"):
            continue
        usual = {}
        for k, n in re.findall(r"(\d+)\s*:\s*(\d+)", str(row[idx["よく使うシフト(参考)"]] or "")):
            usual[int(k)] = int(n)
        days = str(row[idx["勤務可能曜日"]] or "")
        rules[str(name).strip()] = StaffRule(
            target=str(row[idx["対象(○/×)"]] or "").strip() in ("○", "〇", "o", "O", "1", "TRUE", "True"),
            allowed_shifts=_parse_list(row[idx["勤務可能シフト"]]),
            allowed_days="".join(c for c in WEEKDAYS + HOLIDAY if c in days) if days else WEEKDAYS + HOLIDAY,
            target_days=num(row, "目標出勤日数"),
            max_days=num(row, "最大出勤日数"),
            max_hours=num(row, "最大時間/月", float),
            max_consecutive=num(row, "最大連勤"),
            off_mark="空欄" if "空" in str(row[idx["休日の書き方(0/空欄)"]] or "") else "0",
            usual=usual,
        )

    glob = dict(DEFAULT_GLOBAL)
    if SETTINGS_GLOBAL in wb.sheetnames:
        for k, v, *_ in wb[SETTINGS_GLOBAL].iter_rows(min_row=2, values_only=True):
            if k is not None and v not in (None, ""):
                glob[str(k).strip()] = v
    return rules, glob


# ==============================
# 最適化
# ==============================
@dataclass
class Result:
    status: str
    assignment: dict[tuple[int, int], int | None]   # (スタッフ番号, 日) → シフトNO. / 0(休)
    report: list[str]
    shortage_rows: list[list]
    staff_rows: list[list]


def solve(book: Workbook, rules: dict[str, StaffRule], glob: dict, clear: bool = False) -> Result:
    days = range(1, book.num_days + 1)
    timed = {k: sh for k, sh in book.shifts.items() if sh.start is not None}
    w_short = int(glob["重み:人手不足(15分あたり)"])
    w_over = int(glob["重み:人手過剰(15分あたり)"])
    w_days = int(glob["重み:目標日数とのずれ(1日あたり)"])
    w_pref = int(glob["重み:いつもと違うシフト"])
    default_consec = int(glob["最大連勤(既定)"])

    def pattern(day: int) -> dict[int, float]:
        name = str(glob.get(book.day_category(day), ""))
        if name not in book.demand:
            raise ValueError(f"モデル時間にパターン「{name}」がありません（全体設定の「{book.day_category(day)}」）")
        return book.demand[name]

    model = cp_model.CpModel()
    x: dict[tuple[int, int, int], cp_model.IntVar] = {}
    work: dict[tuple[int, int], cp_model.LinearExpr] = {}
    fixed_cover = {(d, h): 0 for d in days for h in range(MODEL_FIRST_HOUR, 31)}
    penalties = []
    targets = []

    for si, s in enumerate(book.staff):
        rule = rules.get(s.name, StaffRule(target=False))
        if not rule.target:
            # AIは触らないが、入っているシフトは人数に数える
            for d in days:
                v = s.grid[d - 1]
                if v in timed:
                    for h in range(MODEL_FIRST_HOUR, 31):
                        fixed_cover[d, h] += timed[v].minutes_in_hour(h) // 15
            continue
        targets.append(si)
        allowed = [k for k in (rule.allowed_shifts or list(timed)) if k in timed]
        usual_max = max(rule.usual.values(), default=0)
        for d in days:
            req = None if clear else s.grid[d - 1]
            opts = set(allowed)
            if req and req in book.shifts:
                opts = {req}
            elif req is None and book.day_category(d) not in rule.allowed_days:
                opts = set()
            if req == 0:
                opts = set()
            for k in opts:
                x[si, d, k] = model.NewBoolVar(f"x_{si}_{d}_{k}")
                if req is None and usual_max and w_pref:
                    cost = round(w_pref * (1 - rule.usual.get(k, 0) / usual_max))
                    if cost:
                        penalties.append(cost * x[si, d, k])
            vs = [x[si, d, k] for k in opts]
            if vs:
                model.Add(sum(vs) <= 1)
                if req:
                    model.Add(sum(vs) == 1)
            work[si, d] = sum(vs) if vs else 0

        n_days = sum(work[si, d] for d in days)
        if isinstance(n_days, int):
            continue                      # 勤務できる日がない
        if rule.max_days is not None:
            model.Add(n_days <= rule.max_days)
        if rule.max_hours is not None:
            quarters = sum(round(book.shifts[k].hours * 4) * v for (i, _, k), v in x.items() if i == si)
            model.Add(quarters <= round(rule.max_hours * 4))
        consec = rule.max_consecutive or default_consec
        for d0 in range(1, book.num_days - consec + 1):
            window = [work[si, d] for d in range(d0, d0 + consec + 1)]
            model.Add(sum(window) <= consec)
        if rule.target_days is not None:
            dev = model.NewIntVar(0, book.num_days, f"dev_{si}")
            model.AddAbsEquality(dev, n_days - rule.target_days)
            penalties.append(w_days * dev)

    # 時間帯ごとの必要人時（15分単位）との差
    by_day: dict[int, list] = {d: [] for d in days}
    for (_, d, k), v in x.items():
        if k in timed:
            by_day[d].append((timed[k], v))
    short_vars = {}
    for d in days:
        dem = pattern(d)
        for h in range(MODEL_FIRST_HOUR, 31):
            need = round(dem.get(h, 0) * 4)
            terms = [sh.minutes_in_hour(h) // 15 * v for sh, v in by_day[d] if sh.minutes_in_hour(h)]
            if not terms and not need:
                continue
            cover = sum(terms) + fixed_cover[d, h]
            short = model.NewIntVar(0, 4 * 50, f"short_{d}_{h}")
            over = model.NewIntVar(0, 4 * 50, f"over_{d}_{h}")
            model.Add(cover - need == over - short)
            penalties += [w_short * short, w_over * over]
            short_vars[d, h] = (short, need)

    model.Minimize(sum(penalties))
    solver = cp_model.CpSolver()
    solver.parameters.max_time_in_seconds = float(glob["計算時間(秒)"])
    solver.parameters.num_workers = 8
    status = solver.Solve(model)
    status_name = {cp_model.OPTIMAL: "最適解", cp_model.FEASIBLE: "実行可能解（時間切れ。計算時間を延ばすと改善の余地あり）"}.get(status)
    if status_name is None:
        return Result("解なし（確定シフト・休み希望・上限の設定が両立しません）", {}, [], [], [])

    assignment: dict[tuple[int, int], int | None] = {}
    for si in targets:
        rule = rules[book.staff[si].name]
        for d in days:
            k = next((k for (i, dd, k), v in x.items() if i == si and dd == d and solver.Value(v)), None)
            assignment[si, d] = k if k is not None else (0 if rule.off_mark == "0" else None)

    # --- レポート ---
    lines = [f"結果: {status_name}", ""]
    shortage_rows = []
    total_short = 0
    for d in days:
        slots = []
        for h in range(MODEL_FIRST_HOUR, 31):
            if (d, h) in short_vars:
                q = solver.Value(short_vars[d, h][0])
                if q:
                    slots.append(f"{h}時台 -{q / 4:g}h")
                    total_short += q
                    shortage_rows.append([f"{book.month}/{d}", book.weekday(d), f"{h}時台",
                                          short_vars[d, h][1] / 4, q / 4])
        if slots:
            lines.append(f"  {book.month}/{d}({book.weekday(d)}) 不足: " + ", ".join(slots))
    lines.insert(1, f"人手不足の合計: {total_short / 4:g} 人時" + ("（不足なし）" if not total_short else ""))

    staff_rows = []
    lines += ["", "スタッフ別:"]
    for si in targets:
        s, rule = book.staff[si], rules[book.staff[si].name]
        ks = [assignment[si, d] for d in days if assignment[si, d]]
        hrs = sum(book.shifts[k].hours for k in ks)
        staff_rows.append([s.name, len(ks), rule.target_days, hrs, rule.max_hours])
        lines.append(f"  {s.name}: {len(ks)}日 {hrs:g}時間"
                     + (f"（目標{rule.target_days}日）" if rule.target_days is not None else ""))
    return Result(status_name, assignment, lines, shortage_rows, staff_rows)


# ==============================
# 書き込み（xlsmのマクロ・画像・ボタンを壊さないよう、シートのXMLだけ書き換える）
# ==============================
def _col_letter(c: int) -> str:
    s = ""
    while c:
        c, r = divmod(c - 1, 26)
        s = chr(65 + r) + s
    return s


def _col_index(letters: str) -> int:
    n = 0
    for ch in letters:
        n = n * 26 + ord(ch) - 64
    return n


def _sheet_xml_path(z: zipfile.ZipFile, sheet_name: str) -> str:
    wbxml = z.read("xl/workbook.xml").decode("utf8")
    rels = z.read("xl/_rels/workbook.xml.rels").decode("utf8")
    for m in re.finditer(r"<sheet\b[^>]*>", wbxml):
        tag = m.group(0)
        name = re.search(r'name="([^"]*)"', tag).group(1)
        if name.replace("&amp;", "&") == sheet_name:
            rid = re.search(r'r:id="([^"]*)"', tag).group(1)
            for rm in re.finditer(r"<Relationship\b[^>]*>", rels):
                if f'Id="{rid}"' in rm.group(0):
                    target = re.search(r'Target="([^"]*)"', rm.group(0)).group(1)
                    return target.lstrip("/") if target.startswith("/") else "xl/" + target
    raise KeyError(sheet_name)


def _set_cells(xml: str, values: dict[str, int | None]) -> str:
    """values: {"E5": 1, "F5": None(空欄)} をシートXMLに書き込む"""
    by_row: dict[int, dict[str, int | None]] = {}
    for ref, v in values.items():
        by_row.setdefault(int(re.search(r"\d+", ref).group(0)), {})[ref] = v

    def cell_xml(ref, style, v):
        s = f' s="{style}"' if style else ""
        return f'<c r="{ref}"{s}/>' if v is None else f'<c r="{ref}"{s}><v>{v}</v></c>'

    def fix_row(m: re.Match) -> str:
        row_xml = m.group(0)
        rnum = int(m.group(1))
        todo = by_row.get(rnum)
        if not todo:
            return row_xml
        if row_xml.endswith("/>"):
            row_xml = row_xml[:-2] + "></row>"
        head_end = row_xml.index(">") + 1
        head, body = row_xml[:head_end], row_xml[head_end:-len("</row>")]
        cells = re.findall(r'<c r="([A-Z]+)\d+"[^>]*?(?:/>|>.*?</c>)', body, flags=re.S)
        parts = re.findall(r'<c r="[A-Z]+\d+"[^>]*?(?:/>|>.*?</c>)', body, flags=re.S)
        cur = dict(zip(cells, parts))
        for ref, v in todo.items():
            col = re.match(r"[A-Z]+", ref).group(0)
            old = cur.get(col)
            style = re.search(r'\ss="(\d+)"', old).group(1) if old and re.search(r'\ss="(\d+)"', old) else None
            cur[col] = cell_xml(ref, style, v)
        rest = re.sub(r'<c r="[A-Z]+\d+"[^>]*?(?:/>|>.*?</c>)', "", body, flags=re.S)
        ordered = "".join(cur[c] for c in sorted(cur, key=_col_index))
        return head + ordered + rest + "</row>"

    return re.sub(r'<row r="(\d+)"[^>]*?(?:/>|>.*?</row>)', fix_row, xml, flags=re.S)


def write_assignment(book: Workbook, result: Result, out_path: str):
    values = {}
    for (si, d), k in result.assignment.items():
        values[f"{_col_letter(PERSONAL_DAY1_COL + d - 1)}{book.staff[si].row}"] = k

    with zipfile.ZipFile(book.path) as zin:
        target = _sheet_xml_path(zin, SHEET_PERSONAL)
        with zipfile.ZipFile(out_path, "w") as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename == target:
                    data = _set_cells(data.decode("utf8"), values).encode("utf8")
                elif item.filename == "xl/workbook.xml":
                    # 開いたときに数式（日別シート・合計）を再計算させる
                    xml = data.decode("utf8")
                    if "<calcPr" in xml:
                        xml = re.sub(r"<calcPr\b([^>]*?)\s*/>",
                                     lambda m: "<calcPr" + re.sub(r'\sfullCalcOnLoad="[^"]*"', "", m.group(1))
                                     + ' fullCalcOnLoad="1"/>', xml, count=1)
                    data = xml.encode("utf8")
                zout.writestr(item, data)


def save_report(path: str, book: Workbook, result: Result):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "人手不足"
    ws.append(["日付", "曜日", "時間帯", "必要人時", "不足人時"])
    for r in result.shortage_rows:
        ws.append(r)
    ws2 = wb.create_sheet("スタッフ別")
    ws2.append(["氏名", "出勤日数", "目標日数", "時間", "上限時間"])
    for r in result.staff_rows:
        ws2.append(r)
    wb.save(path)


# ==============================
# コマンド
# ==============================
def cmd_learn(args):
    book = load_workbook(args.excel)
    rules = learn(book)
    save_settings(args.settings, book, rules)
    print(f"{book.year}年{book.month}月のシフトから学習しました: {args.settings}")
    for s in book.staff:
        r = rules[s.name]
        print(f"  {'○' if r.target else '×'} {s.name}: シフト[{_fmt_list(r.allowed_shifts)}] "
              f"曜日[{r.allowed_days}] {r.target_days}日")


def cmd_generate(args):
    book = load_workbook(args.excel)
    if args.settings and os.path.exists(args.settings):
        rules, glob = load_settings(args.settings)
    else:
        print("設定ファイルがないので、このファイルの入力内容から学習して作成します")
        rules, glob = learn(book), dict(DEFAULT_GLOBAL)
    if args.time:
        glob["計算時間(秒)"] = args.time
    print(f"{book.year}年{book.month}月 / スタッフ{len(book.staff)}名 / シフト{len(book.shifts)}種類 / "
          f"{book.num_days}日間 / 祝日{sorted(book.holidays)}")
    result = solve(book, rules, glob, clear=args.clear)
    print("\n".join(result.report) or result.status)
    if not result.assignment:
        sys.exit(1)
    out = args.output or re.sub(r"(\.xls[xm])$", r"_AI\1", os.path.basename(args.excel))
    write_assignment(book, result, out)
    print(f"\n保存しました: {out}")
    if args.report:
        save_report(args.report, book, result)
        print(f"レポート: {args.report}")
    return out


def main(argv=None):
    p = argparse.ArgumentParser(description="AIシフト作成ツール")
    sub = p.add_subparsers(dest="cmd", required=True)
    a = sub.add_parser("learn", help="完成済みの月から勤務傾向を学習して設定ファイルを作る")
    a.add_argument("excel")
    a.add_argument("-s", "--settings", default="AI設定.xlsx")
    a.set_defaults(func=cmd_learn)
    g = sub.add_parser("generate", help="シフトを自動作成する")
    g.add_argument("excel")
    g.add_argument("-s", "--settings", default="AI設定.xlsx")
    g.add_argument("-o", "--output")
    g.add_argument("-r", "--report", help="レポートを保存する .xlsx")
    g.add_argument("-t", "--time", type=int, help="計算時間(秒)")
    g.add_argument("--clear", action="store_true", help="入力済みのシフトを無視して全部作り直す")
    g.set_defaults(func=cmd_generate)
    args = p.parse_args(argv)
    return args.func(args)


if __name__ == "__main__":
    main()
